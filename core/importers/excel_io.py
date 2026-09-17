import openpyxl
from openpyxl.utils import get_column_letter
from django.http import HttpResponse

MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5MB
MAX_UPLOAD_ROWS = 5000


def build_template_xlsx(spec):
    """Builds a downloadable .xlsx: a header row plus one example row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec.label[:31]  # Excel sheet names cap at 31 characters

    for col_idx, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col.header)
        cell.font = openpyxl.styles.Font(bold=True)
        ws.cell(row=2, column=col_idx, value=col.example)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(18, len(col.header) + 4)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{spec.key}_template.xlsx"'
    wb.save(response)
    return response


class UploadTooLarge(Exception):
    pass


def read_uploaded_rows(uploaded_file, spec):
    """
    Reads an uploaded .xlsx matched by header name, not column position.
    Returns: (list_of_dict_rows, list_of_missing_headers)
    Each dict row has 'row_number' and 'data' (dict of column_key->value).
    """
    if uploaded_file.size > MAX_UPLOAD_BYTES:
        raise UploadTooLarge(f"File is too large (max {MAX_UPLOAD_BYTES // 1024 // 1024}MB).")

    wb = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []

    # Map header names (lowercase, stripped) to column indices
    header_positions = {
        str(h).strip().lower(): idx
        for idx, h in enumerate(header_row) if h is not None
    }

    col_to_position = {}
    for col in spec.columns:
        pos = header_positions.get(col.header.strip().lower())
        col_to_position[col.key] = pos  # None if missing

    missing_headers = [c.header for c in spec.columns if col_to_position.get(c.key) is None]

    results = []
    for row_number, raw_row in enumerate(rows_iter, start=2):
        if row_number - 1 > MAX_UPLOAD_ROWS:
            raise UploadTooLarge(f"File has more than {MAX_UPLOAD_ROWS} data rows.")
        if raw_row is None or not any(raw_row):
            continue  # blank row, skip

        row_data = {}
        for col in spec.columns:
            pos = col_to_position.get(col.key)
            value = raw_row[pos] if pos is not None and pos < len(raw_row) else None
            row_data[col.key] = str(value).strip() if value is not None else ''

        results.append({'row_number': row_number, 'data': row_data})

    return results, missing_headers


def _sanitize_for_export(value):
    """Prevent formula injection by prefixing =, +, -, @ with a single quote."""
    text = str(value) if value is not None else ''
    if text and text[:1] in ('=', '+', '-', '@'):
        return "'" + text
    return text


def build_export_xlsx(spec):
    """Builds an .xlsx of all data for this spec, using same column headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec.label[:31]

    for col_idx, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col.header)
        cell.font = openpyxl.styles.Font(bold=True)

    for row_idx, instance in enumerate(spec.export_queryset(), start=2):
        row_dict = spec.row_from_instance(instance)
        for col_idx, col in enumerate(spec.columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_sanitize_for_export(row_dict.get(col.key, '')))

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{spec.key}_export.xlsx"'
    wb.save(response)
    return response