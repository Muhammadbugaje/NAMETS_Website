"""
Governance app — Django admin registration.

Provides full CRUD + bulk actions + custom filters for all governance models.
Superusers and the ICT Head use this as their direct management surface;
the custom EXCO dashboards (built later) are the primary UX for other offices.
"""

from django.contrib import admin
from django.contrib import messages
from django.db.models import Count, Q, F
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import path, reverse
from django.utils import timezone
from django.utils.html import format_html, mark_safe
from unfold.admin import ModelAdmin, TabularInline

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    AuditLog, Task, Proposal, Vote,
    NominationIntake, SelectionRecord, OathAcknowledgment,
    HandoverLog, GraduationProfile,
)


# ============================================================
# CUSTOM FILTERS
# ============================================================

class TaskOverdueFilter(admin.SimpleListFilter):
    title = 'deadline'
    parameter_name = 'deadline_state'

    def lookups(self, request, model_admin):
        return [
            ('overdue', 'Overdue'),
            ('today', 'Due today'),
            ('week', 'Due this week'),
            ('none', 'No deadline set'),
        ]

    def queryset(self, request, queryset):
        now = timezone.now()
        today_end = now.replace(hour=23, minute=59, second=59)
        week_end = now + timezone.timedelta(days=7)
        v = self.value()
        if v == 'overdue':
            return queryset.filter(
                deadline__lt=now,
                status__in=['pending', 'in_progress'],
            )
        if v == 'today':
            return queryset.filter(deadline__date=now.date())
        if v == 'week':
            return queryset.filter(deadline__gte=now, deadline__lte=week_end)
        if v == 'none':
            return queryset.filter(deadline__isnull=True)
        return queryset


class AuditActionFilter(admin.SimpleListFilter):
    title = 'action group'
    parameter_name = 'action_group'

    def lookups(self, request, model_admin):
        return [
            ('permissions', 'Permission changes'),
            ('votes', 'Voting activity'),
            ('handover', 'Handover / dissolve'),
            ('finance', 'Payments'),
            ('imports', 'Bulk imports/exports'),
            ('auth', 'Logins / logouts'),
        ]

    def queryset(self, request, queryset):
        v = self.value()
        if v == 'permissions':
            return queryset.filter(action__in=['permission_granted', 'permission_revoked'])
        if v == 'votes':
            return queryset.filter(action__in=['vote_cast', 'proposal_opened', 'proposal_closed'])
        if v == 'handover':
            return queryset.filter(action='handover')
        if v == 'finance':
            return queryset.filter(action='payment_logged')
        if v == 'imports':
            return queryset.filter(action__in=['bulk_import', 'bulk_export'])
        if v == 'auth':
            return queryset.filter(action__in=['login', 'logout'])
        return queryset


# ============================================================
# AUDIT LOG
# ============================================================

@admin.register(AuditLog)
class AuditLogAdmin(ModelAdmin):
    list_display = (
        'timestamp_display', 'actor_display_col', 'action_badge',
        'office', 'object_repr_short', 'ip_address',
    )
    list_filter = ('action', AuditActionFilter, 'office', 'timestamp')
    search_fields = ('object_repr', 'target_model', 'user__email', 'user__first_name', 'user__last_name')
    readonly_fields = (
        'user', 'action', 'office', 'target_model', 'target_id',
        'object_repr', 'changes_pretty', 'ip_address', 'user_agent',
        'timestamp',
    )
    date_hierarchy = 'timestamp'
    list_per_page = 50
    ordering = ('-timestamp',)

    fieldsets = (
        ('Who', {'fields': ('user', 'office', 'ip_address', 'user_agent')}),
        ('What', {'fields': ('action', 'object_repr', 'target_model', 'target_id')}),
        ('Details', {'fields': ('changes_pretty',)}),
        ('When', {'fields': ('timestamp',)}),
    )

    actions = ['export_selected_to_excel']

    # ---------- Displays ----------

    def timestamp_display(self, obj):
        return format_html(
            '<span style="color:#0F3D2E;font-weight:600;">{}</span>'
            '<br><span style="color:#8a9b95;font-size:0.8em;">{}</span>',
            obj.timestamp.strftime('%Y-%m-%d'),
            obj.timestamp.strftime('%H:%M:%S'),
        )
    timestamp_display.short_description = 'When'
    timestamp_display.admin_order_field = 'timestamp'

    def actor_display_col(self, obj):
        if not obj.user:
            return mark_safe('<span style="color:#8a9b95;">System</span>')
        return format_html(
            '<strong>{}</strong><br>'
            '<span style="color:#8a9b95;font-size:0.8em;">{}</span>',
            obj.user.get_full_name() or obj.user.username,
            obj.user.email or '',
        )
    actor_display_col.short_description = 'Actor'

    def action_badge(self, obj):
        colors = {
            'create': ('#e8f2ec', '#1a6b3c'),
            'update': ('#e6f0fb', '#185fa5'),
            'delete': ('#fdecea', '#9b2c2c'),
            'permission_granted': ('#f0eafb', '#6b21a8'),
            'permission_revoked': ('#fef3e8', '#a8671a'),
            'vote_cast': ('#f5e9c8', '#8a6a1a'),
            'proposal_opened': ('#e0f7f4', '#0d7a6b'),
            'proposal_closed': ('#e0f7f4', '#0d7a6b'),
            'handover': ('#fdecea', '#9b2c2c'),
            'payment_logged': ('#e8f2ec', '#1a6b3c'),
            'bulk_import': ('#f0eafb', '#6b21a8'),
            'bulk_export': ('#f0eafb', '#6b21a8'),
            'login': ('#f0f4f2', '#4a5568'),
            'logout': ('#f0f4f2', '#4a5568'),
        }
        bg, fg = colors.get(obj.action, ('#f0f4f2', '#4a5568'))
        return format_html(
            '<span style="background:{};color:{};padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">{}</span>',
            bg, fg, obj.get_action_display()
        )
    action_badge.short_description = 'Action'
    action_badge.admin_order_field = 'action'

    def object_repr_short(self, obj):
        text = obj.object_repr or '—'
        if len(text) > 60:
            text = text[:60] + '…'
        return text
    object_repr_short.short_description = 'Object'

    def changes_pretty(self, obj):
        import json
        if not obj.changes:
            return mark_safe('<span style="color:#8a9b95;">No details</span>')
        try:
            pretty = json.dumps(obj.changes, indent=2, default=str)
        except Exception:
            pretty = str(obj.changes)
        return format_html(
            '<pre style="background:#f8faf9;padding:12px;border-radius:8px;'
            'font-size:0.85em;line-height:1.5;max-height:400px;overflow:auto;">{}</pre>',
            pretty
        )
    changes_pretty.short_description = 'Details'

    # ---------- Bulk actions ----------

    @admin.action(description='Export selected log entries to Excel')
    def export_selected_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Audit Log'

        headers = ['Timestamp', 'Actor', 'Email', 'Action', 'Office',
                   'Target Model', 'Object', 'IP Address', 'Details']
        gold = PatternFill(start_color='C9A84C', end_color='C9A84C', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for i, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.fill = gold
            cell.font = header_font

        for i, log in enumerate(queryset, start=2):
            ws.cell(row=i, column=1, value=log.timestamp.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=i, column=2, value=log.actor_display)
            ws.cell(row=i, column=3, value=log.user.email if log.user else '')
            ws.cell(row=i, column=4, value=log.get_action_display())
            ws.cell(row=i, column=5, value=log.office.name if log.office else '')
            ws.cell(row=i, column=6, value=log.target_model or '')
            ws.cell(row=i, column=7, value=log.object_repr or '')
            ws.cell(row=i, column=8, value=log.ip_address or '')
            ws.cell(row=i, column=9, value=str(log.changes or ''))

        for i, w in enumerate([20, 25, 25, 20, 20, 20, 40, 15, 50], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"audit_log_{timezone.now():%Y%m%d_%H%M}.xlsx"
        response['Content-Disposition'] = f'attachment; filename={filename}'
        wb.save(response)
        self.message_user(request, f"{queryset.count()} log entries exported.", messages.SUCCESS)
        return response


# ============================================================
# TASKS
# ============================================================

@admin.register(Task)
class TaskAdmin(ModelAdmin):
    list_display = (
        'title_short', 'assignee_col', 'priority_badge', 'status_badge',
        'deadline_col', 'overdue_indicator', 'created_at',
    )
    list_filter = ('status', 'priority', TaskOverdueFilter, 'assigned_office')
    search_fields = ('title', 'description', 'assigned_to__email', 'assigned_to__first_name', 'assigned_to__last_name')
    readonly_fields = ('created_at', 'updated_at', 'completed_at', 'completed_by')
    date_hierarchy = 'created_at'
    list_per_page = 40
    ordering = ('status', 'deadline', '-created_at')

    fieldsets = (
        ('Task', {'fields': ('title', 'description', 'priority', 'deadline')}),
        ('Assignment', {
            'fields': ('assigned_to', 'assigned_office', 'assigned_by'),
            'description': 'Assign to a person OR an office — not both.',
        }),
        ('Progress', {'fields': ('status', 'completion_proof')}),
        ('Completion', {
            'fields': ('completed_at', 'completed_by'),
            'classes': ('collapse',),
        }),
        ('Link (optional)', {
            'fields': ('related_object_type', 'related_object_id'),
            'classes': ('collapse',),
        }),
        ('Meta', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',),
        }),
    )

    actions = [
        'mark_completed', 'mark_in_progress', 'mark_cancelled',
        'mark_urgent', 'mark_high_priority', 'export_to_excel',
    ]

    # ---------- Displays ----------

    def title_short(self, obj):
        return format_html('<strong style="color:#0F3D2E;">{}</strong>', obj.title[:70])
    title_short.short_description = 'Task'
    title_short.admin_order_field = 'title'

    def assignee_col(self, obj):
        if obj.assigned_to:
            return format_html(
                '<span style="color:#185fa5;">👤 {}</span>',
                obj.assignee_display
            )
        if obj.assigned_office:
            return format_html(
                '<span style="color:#6b21a8;">🏛️ {}</span>',
                obj.assignee_display
            )
        return mark_safe('<span style="color:#c0392b;">Unassigned</span>')
    assignee_col.short_description = 'Assigned to'

    def priority_badge(self, obj):
        colors = {
            'low': ('#e8f2ec', '#1a6b3c', 'Low'),
            'medium': ('#f5e9c8', '#8a6a1a', 'Medium'),
            'high': ('#fef3e8', '#a8671a', 'High'),
            'urgent': ('#fdecea', '#9b2c2c', 'Urgent'),
        }
        bg, fg, label = colors.get(obj.priority, ('#f0f4f2', '#4a5568', obj.priority))
        return format_html(
            '<span style="background:{};color:{};padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">{}</span>',
            bg, fg, label
        )
    priority_badge.short_description = 'Priority'
    priority_badge.admin_order_field = 'priority'

    def status_badge(self, obj):
        if obj.is_overdue:
            return mark_safe(
                '<span style="background:#fdecea;color:#9b2c2c;padding:2px 9px;'
                'border-radius:10px;font-size:0.72em;font-weight:700;">Overdue</span>'
            )
        colors = {
            'pending': ('#f5e9c8', '#8a6a1a'),
            'in_progress': ('#e6f0fb', '#185fa5'),
            'completed': ('#e8f2ec', '#1a6b3c'),
            'cancelled': ('#f0f0f0', '#777'),
        }
        bg, fg = colors.get(obj.status, ('#f0f4f2', '#4a5568'))
        return format_html(
            '<span style="background:{};color:{};padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">{}</span>',
            bg, fg, obj.get_status_display()
        )
    status_badge.short_description = 'Status'
    status_badge.admin_order_field = 'status'

    def deadline_col(self, obj):
        if not obj.deadline:
            return mark_safe('<span style="color:#8a9b95;">—</span>')
        return obj.deadline.strftime('%b %d, %Y · %H:%M')
    deadline_col.short_description = 'Deadline'
    deadline_col.admin_order_field = 'deadline'

    def overdue_indicator(self, obj):
        if obj.is_overdue:
            return mark_safe('<span style="color:#c0392b;font-weight:800;">⚠️</span>')
        return ''
    overdue_indicator.short_description = ''

    # ---------- Bulk actions ----------

    @admin.action(description='✅ Mark as completed')
    def mark_completed(self, request, queryset):
        count = 0
        for t in queryset:
            if t.status != 'completed':
                t.mark_completed(request.user)
                count += 1
        self.message_user(request, f"{count} task(s) marked completed.", messages.SUCCESS)

    @admin.action(description='🔄 Mark as in progress')
    def mark_in_progress(self, request, queryset):
        n = queryset.update(status='in_progress', updated_at=timezone.now())
        self.message_user(request, f"{n} task(s) set to In Progress.", messages.SUCCESS)

    @admin.action(description='❌ Mark as cancelled')
    def mark_cancelled(self, request, queryset):
        n = queryset.update(status='cancelled', updated_at=timezone.now())
        self.message_user(request, f"{n} task(s) cancelled.", messages.WARNING)

    @admin.action(description='🔴 Set priority → Urgent')
    def mark_urgent(self, request, queryset):
        n = queryset.update(priority='urgent', updated_at=timezone.now())
        self.message_user(request, f"{n} task(s) set to Urgent.", messages.WARNING)

    @admin.action(description='🟠 Set priority → High')
    def mark_high_priority(self, request, queryset):
        n = queryset.update(priority='high', updated_at=timezone.now())
        self.message_user(request, f"{n} task(s) set to High.", messages.SUCCESS)

    @admin.action(description='📥 Export selected tasks to Excel')
    def export_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Tasks'
        headers = ['Title', 'Assigned To', 'Priority', 'Status', 'Deadline',
                   'Created', 'Completed', 'Overdue']
        gold = PatternFill(start_color='C9A84C', end_color='C9A84C', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.fill = gold
            c.font = header_font
        for i, t in enumerate(queryset, start=2):
            ws.cell(row=i, column=1, value=t.title)
            ws.cell(row=i, column=2, value=t.assignee_display)
            ws.cell(row=i, column=3, value=t.get_priority_display())
            ws.cell(row=i, column=4, value=t.get_status_display())
            ws.cell(row=i, column=5, value=t.deadline.strftime('%Y-%m-%d %H:%M') if t.deadline else '')
            ws.cell(row=i, column=6, value=t.created_at.strftime('%Y-%m-%d'))
            ws.cell(row=i, column=7, value=t.completed_at.strftime('%Y-%m-%d') if t.completed_at else '')
            ws.cell(row=i, column=8, value='Yes' if t.is_overdue else 'No')
        for i, w in enumerate([35, 25, 12, 15, 22, 15, 15, 10], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"tasks_{timezone.now():%Y%m%d_%H%M}.xlsx"
        response['Content-Disposition'] = f'attachment; filename={filename}'
        wb.save(response)
        self.message_user(request, f"{queryset.count()} task(s) exported.", messages.SUCCESS)
        return response


# ============================================================
# VOTES (inline inside Proposal)
# ============================================================

class VoteInline(TabularInline):
    model = Vote
    extra = 0
    can_delete = False
    fields = ('voter', 'choice', 'feedback_short', 'cast_at')
    readonly_fields = ('voter', 'choice', 'feedback_short', 'cast_at')
    per_page = 25
    classes = ('collapse',)

    def feedback_short(self, obj):
        if not obj.feedback:
            return '—'
        return obj.feedback[:80] + ('…' if len(obj.feedback) > 80 else '')
    feedback_short.short_description = 'Feedback'

    def has_add_permission(self, request, obj=None):
        return False


# ============================================================
# PROPOSALS
# ============================================================

@admin.register(Proposal)
class ProposalAdmin(ModelAdmin):
    list_display = (
        'title_short', 'proposed_by_col', 'status_badge',
        'voting_window', 'participation_col', 'result_col',
    )
    list_filter = ('status', 'proposed_by', 'voting_start')
    search_fields = ('title', 'description', 'proposed_by__email', 'proposed_by__first_name')
    readonly_fields = (
        'created_at', 'updated_at', 'decided_at',
        'tally_display', 'result_summary',
    )
    inlines = [VoteInline]
    date_hierarchy = 'created_at'
    list_per_page = 30
    ordering = ('-created_at',)

    fieldsets = (
        ('Proposal', {
            'fields': ('title', 'description', 'proposed_budget', 'proposed_by')
        }),
        ('Voting Window', {
            'fields': ('voting_start', 'voting_end', 'status')
        }),
        ('Rules', {
            'fields': ('quorum_percentage', 'approval_percentage'),
            'description': 'Quorum = minimum % of eligible voters who must vote. '
                           'Approval = minimum % of cast votes that must be "Agree".',
        }),
        ('Live Tally', {
            'fields': ('tally_display',),
        }),
        ('Outcome', {
            'fields': ('decided_at', 'result_summary'),
            'classes': ('collapse',),
        }),
        ('Meta', {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',),
        }),
    )

    actions = ['open_voting', 'close_and_decide', 'cancel_proposals', 'export_to_excel']

    # ---------- Displays ----------

    def title_short(self, obj):
        return format_html('<strong style="color:#0F3D2E;">{}</strong>', obj.title[:60])
    title_short.short_description = 'Proposal'
    title_short.admin_order_field = 'title'

    def proposed_by_col(self, obj):
        return obj.proposed_by.get_full_name() or obj.proposed_by.username
    proposed_by_col.short_description = 'Proposed by'

    def status_badge(self, obj):
        colors = {
            'draft': ('#f0f4f2', '#4a5568'),
            'voting': ('#e6f0fb', '#185fa5'),
            'passed': ('#e8f2ec', '#1a6b3c'),
            'rejected': ('#fdecea', '#9b2c2c'),
            'cancelled': ('#f0f0f0', '#777'),
        }
        bg, fg = colors.get(obj.status, ('#f0f4f2', '#4a5568'))
        pulse = ''
        if obj.status == 'voting' and obj.voting_end > timezone.now():
            pulse = ' ●'
        return format_html(
            '<span style="background:{};color:{};padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">{}{}</span>',
            bg, fg, obj.get_status_display(), pulse
        )
    status_badge.short_description = 'Status'
    status_badge.admin_order_field = 'status'

    def voting_window(self, obj):
        return format_html(
            '<span style="font-size:0.85em;">{}<br>→ {}</span>',
            obj.voting_start.strftime('%b %d, %H:%M'),
            obj.voting_end.strftime('%b %d, %H:%M'),
        )
    voting_window.short_description = 'Voting window'

    def participation_col(self, obj):
        pct = obj.participation_percentage
        color = '#1a6b3c' if pct >= obj.quorum_percentage else '#a8671a'
        return format_html(
            '<strong style="color:{};">{}%</strong><br>'
            '<span style="font-size:0.75em;color:#8a9b95;">{}/{} voted</span>',
            color, pct, obj.total_votes_cast, obj.total_eligible_voters
        )
    participation_col.short_description = 'Participation'

    def result_col(self, obj):
        if obj.status in ('passed', 'rejected'):
            return format_html(
                '<span style="font-size:0.8em;">{}</span>',
                (obj.result_summary[:60] + '…') if len(obj.result_summary) > 60 else obj.result_summary
            )
        if obj.status == 'voting':
            return mark_safe(
                '<span style="color:#185fa5;font-weight:600;">In progress…</span>'
            )
        return '—'
    result_col.short_description = 'Result'

    def tally_display(self, obj):
        if not obj.pk:
            return '—'
        t = obj.tally()
        return format_html(
            '<div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(140px,1fr));'
            'gap:0.75rem;background:#f8faf9;padding:1rem;border-radius:12px;">'
            '<div><div style="font-size:0.65rem;text-transform:uppercase;color:#8a9b95;'
            'font-weight:700;letter-spacing:0.05em;">Eligible</div>'
            '<div style="font-size:1.3rem;font-weight:800;color:#0F3D2E;">{eligible}</div></div>'
            '<div><div style="font-size:0.65rem;text-transform:uppercase;color:#8a9b95;'
            'font-weight:700;letter-spacing:0.05em;">Votes cast</div>'
            '<div style="font-size:1.3rem;font-weight:800;color:#0F3D2E;">{cast} '
            '<span style="font-size:0.65em;color:#8a9b95;">({pct}%)</span></div></div>'
            '<div><div style="font-size:0.65rem;text-transform:uppercase;color:#8a9b95;'
            'font-weight:700;letter-spacing:0.05em;">Agree</div>'
            '<div style="font-size:1.3rem;font-weight:800;color:#1a6b3c;">{agree}</div></div>'
            '<div><div style="font-size:0.65rem;text-transform:uppercase;color:#8a9b95;'
            'font-weight:700;letter-spacing:0.05em;">Disagree</div>'
            '<div style="font-size:1.3rem;font-weight:800;color:#9b2c2c;">{disagree}</div></div>'
            '<div><div style="font-size:0.65rem;text-transform:uppercase;color:#8a9b95;'
            'font-weight:700;letter-spacing:0.05em;">Abstain</div>'
            '<div style="font-size:1.3rem;font-weight:800;color:#8a6a1a;">{abstain}</div></div>'
            '</div>',
            eligible=t['eligible'],
            cast=t['cast'],
            pct=t['participation_pct'],
            agree=t['agree'],
            disagree=t['disagree'],
            abstain=t['abstain'],
        )
    tally_display.short_description = 'Current standing'

    # ---------- Bulk actions ----------

    @admin.action(description='🗳️ Open voting')
    def open_voting(self, request, queryset):
        n = 0
        for p in queryset:
            if p.status == 'draft':
                p.status = 'voting'
                p.save(update_fields=['status', 'updated_at'])
                n += 1
        self.message_user(request, f"{n} proposal(s) opened for voting.", messages.SUCCESS)

    @admin.action(description='📊 Close voting & decide result')
    def close_and_decide(self, request, queryset):
        n = 0
        for p in queryset:
            if p.status == 'voting':
                p.decide(by_user=request.user, force=True)
                n += 1
        self.message_user(request, f"{n} proposal(s) decided.", messages.SUCCESS)

    @admin.action(description='🚫 Cancel proposals')
    def cancel_proposals(self, request, queryset):
        n = queryset.update(status='cancelled', updated_at=timezone.now())
        self.message_user(request, f"{n} proposal(s) cancelled.", messages.WARNING)

    @admin.action(description='📥 Export selected proposals to Excel')
    def export_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Proposals'
        headers = ['Title', 'Proposed by', 'Budget', 'Status',
                   'Voting start', 'Voting end', 'Eligible', 'Cast',
                   'Agree', 'Disagree', 'Abstain', 'Result']
        gold = PatternFill(start_color='C9A84C', end_color='C9A84C', fill_type='solid')
        hf = Font(bold=True, color='FFFFFF')
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=h)
            c.fill = gold; c.font = hf
        for i, p in enumerate(queryset, start=2):
            t = p.tally()
            ws.cell(row=i, column=1, value=p.title)
            ws.cell(row=i, column=2, value=p.proposed_by.get_full_name())
            ws.cell(row=i, column=3, value=float(p.proposed_budget))
            ws.cell(row=i, column=4, value=p.get_status_display())
            ws.cell(row=i, column=5, value=p.voting_start.strftime('%Y-%m-%d %H:%M'))
            ws.cell(row=i, column=6, value=p.voting_end.strftime('%Y-%m-%d %H:%M'))
            ws.cell(row=i, column=7, value=t['eligible'])
            ws.cell(row=i, column=8, value=t['cast'])
            ws.cell(row=i, column=9, value=t['agree'])
            ws.cell(row=i, column=10, value=t['disagree'])
            ws.cell(row=i, column=11, value=t['abstain'])
            ws.cell(row=i, column=12, value=p.result_summary)
        for i, w in enumerate([35, 20, 15, 15, 20, 20, 10, 10, 10, 10, 10, 45], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"proposals_{timezone.now():%Y%m%d_%H%M}.xlsx"
        response['Content-Disposition'] = f'attachment; filename={filename}'
        wb.save(response)
        return response


# ============================================================
# VOTES (standalone for search)
# ============================================================

@admin.register(Vote)
class VoteAdmin(ModelAdmin):
    list_display = ('proposal_col', 'voter_col', 'choice_badge', 'cast_at')
    list_filter = ('choice', 'proposal', 'cast_at')
    search_fields = ('voter__email', 'voter__first_name', 'voter__last_name', 'proposal__title')
    readonly_fields = ('proposal', 'voter', 'choice', 'feedback', 'cast_at')
    date_hierarchy = 'cast_at'
    ordering = ('-cast_at',)

    def proposal_col(self, obj):
        return obj.proposal.title[:50]
    proposal_col.short_description = 'Proposal'

    def voter_col(self, obj):
        return obj.voter.get_full_name() or obj.voter.username
    voter_col.short_description = 'Voter'

    def choice_badge(self, obj):
        colors = {
            'agree': ('#e8f2ec', '#1a6b3c', '✓ Agree'),
            'disagree': ('#fdecea', '#9b2c2c', '✗ Disagree'),
            'abstain': ('#f5e9c8', '#8a6a1a', '⊘ Abstain'),
        }
        bg, fg, label = colors.get(obj.choice, ('#f0f4f2', '#4a5568', obj.choice))
        return format_html(
            '<span style="background:{};color:{};padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">{}</span>',
            bg, fg, label
        )
    choice_badge.short_description = 'Choice'


# ============================================================
# SHURA / NAQIB WORKFLOW
# ============================================================

@admin.register(NominationIntake)
class NominationIntakeAdmin(ModelAdmin):
    list_display = (
        'nominated_name', 'office_col', 'nominated_by_col',
        'status_badge', 'session_label', 'created_at',
    )
    list_filter = ('status', 'session_label', 'nominated_office')
    search_fields = ('nominated_name', 'nominated_email', 'nominated_phone', 'reason')
    date_hierarchy = 'created_at'
    list_per_page = 40

    fieldsets = (
        ('Nominee', {
            'fields': ('nominated_name', 'nominated_email', 'nominated_phone')
        }),
        ('Office', {'fields': ('nominated_office', 'session_label')}),
        ('Nomination', {
            'fields': ('nominated_by', 'reason', 'status')
        }),
    )

    actions = ['shortlist', 'accept', 'reject', 'export_to_excel']

    def office_col(self, obj):
        if not obj.nominated_office:
            return mark_safe('<span style="color:#8a9b95;">—</span>')
        return obj.nominated_office.name
    office_col.short_description = 'Office'

    def nominated_by_col(self, obj):
        if not obj.nominated_by:
            return '—'
        return obj.nominated_by.get_full_name() or obj.nominated_by.username
    nominated_by_col.short_description = 'Nominated by'

    def status_badge(self, obj):
        colors = {
            'pending': ('#fef3e8', '#a8671a'),
            'shortlisted': ('#e6f0fb', '#185fa5'),
            'accepted': ('#e8f2ec', '#1a6b3c'),
            'rejected': ('#fdecea', '#9b2c2c'),
        }
        bg, fg = colors.get(obj.status, ('#f0f4f2', '#4a5568'))
        return format_html(
            '<span style="background:{};color:{};padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">{}</span>',
            bg, fg, obj.get_status_display()
        )
    status_badge.short_description = 'Status'

    @admin.action(description='⭐ Mark as Shortlisted')
    def shortlist(self, request, queryset):
        n = queryset.update(status='shortlisted')
        self.message_user(request, f"{n} nomination(s) shortlisted.", messages.SUCCESS)

    @admin.action(description='✅ Mark as Accepted')
    def accept(self, request, queryset):
        n = queryset.update(status='accepted')
        self.message_user(request, f"{n} nomination(s) accepted.", messages.SUCCESS)

    @admin.action(description='❌ Mark as Rejected')
    def reject(self, request, queryset):
        n = queryset.update(status='rejected')
        self.message_user(request, f"{n} nomination(s) rejected.", messages.WARNING)

    @admin.action(description='📥 Export nominations to Excel')
    def export_to_excel(self, request, queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Nominations'
        headers = ['Name', 'Email', 'Phone', 'Office', 'Session',
                   'Nominated by', 'Status', 'Reason', 'Created']
        gold = PatternFill(start_color='C9A84C', end_color='C9A84C', fill_type='solid')
        hf = Font(bold=True, color='FFFFFF')
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=i, value=h); c.fill = gold; c.font = hf
        for i, n in enumerate(queryset, start=2):
            ws.cell(row=i, column=1, value=n.nominated_name)
            ws.cell(row=i, column=2, value=n.nominated_email)
            ws.cell(row=i, column=3, value=n.nominated_phone)
            ws.cell(row=i, column=4, value=n.nominated_office.name if n.nominated_office else '')
            ws.cell(row=i, column=5, value=n.session_label)
            ws.cell(row=i, column=6, value=n.nominated_by.get_full_name() if n.nominated_by else '')
            ws.cell(row=i, column=7, value=n.get_status_display())
            ws.cell(row=i, column=8, value=n.reason)
            ws.cell(row=i, column=9, value=n.created_at.strftime('%Y-%m-%d'))
        for i, w in enumerate([25, 25, 15, 25, 15, 20, 15, 40, 15], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = (
            f'attachment; filename=nominations_{timezone.now():%Y%m%d_%H%M}.xlsx'
        )
        wb.save(response)
        return response


@admin.register(SelectionRecord)
class SelectionRecordAdmin(ModelAdmin):
    list_display = ('session_label', 'finalized_by_col', 'finalized_at', 'published_badge', 'count_display')
    list_filter = ('is_published',)
    search_fields = ('session_label', 'notes')
    readonly_fields = ('finalized_at',)

    fieldsets = (
        ('Session', {
            'fields': ('session_label', 'is_published')
        }),
        ('Finalized by', {
            'fields': ('finalized_by', 'finalized_at')
        }),
        ('Snapshot', {
            'fields': ('snapshot',),
            'description': 'List of {"office": ..., "person_name": ..., "person_email": ...} dicts.',
        }),
        ('Notes', {'fields': ('notes',)}),
    )

    def finalized_by_col(self, obj):
        if not obj.finalized_by:
            return '—'
        return obj.finalized_by.get_full_name() or obj.finalized_by.username
    finalized_by_col.short_description = 'Finalized by'

    def published_badge(self, obj):
        if obj.is_published:
            return mark_safe(
                '<span style="background:#e8f2ec;color:#1a6b3c;padding:2px 9px;'
                'border-radius:10px;font-size:0.72em;font-weight:700;">Published</span>'
            )
        return mark_safe(
            '<span style="background:#f0f0f0;color:#777;padding:2px 9px;'
            'border-radius:10px;font-size:0.72em;font-weight:700;">Draft</span>'
        )
    published_badge.short_description = 'Status'

    def count_display(self, obj):
        return f"{len(obj.snapshot or [])} entries"
    count_display.short_description = 'Entries'


@admin.register(OathAcknowledgment)
class OathAcknowledgmentAdmin(ModelAdmin):
    list_display = ('user_col', 'session_label', 'acknowledged_at', 'notes')
    list_filter = ('session_label', 'acknowledged_at')
    search_fields = ('user__email', 'user__first_name', 'user__last_name')
    readonly_fields = ('acknowledged_at',)

    def user_col(self, obj):
        return obj.user.get_full_name() or obj.user.username
    user_col.short_description = 'User'


# ============================================================
# HANDOVER
# ============================================================

@admin.register(HandoverLog)
class HandoverLogAdmin(ModelAdmin):
    list_display = (
        'timestamp', 'initiated_by_col', 'session_ending',
        'session_starting', 'dissolved_count', 'protected_skipped',
    )
    list_filter = ('timestamp', 'initiated_by')
    readonly_fields = (
        'initiated_by', 'session_ending', 'session_starting',
        'dissolved_count', 'protected_skipped', 'notes', 'timestamp',
    )
    ordering = ('-timestamp',)

    def initiated_by_col(self, obj):
        if not obj.initiated_by:
            return '—'
        return obj.initiated_by.get_full_name() or obj.initiated_by.username
    initiated_by_col.short_description = 'Initiated by'


@admin.register(GraduationProfile)
class GraduationProfileAdmin(ModelAdmin):
    list_display = ('user_col', 'department_name', 'graduation_year', 'offices_count', 'tasks_completed', 'proposals_passed')
    list_filter = ('graduation_year',)
    search_fields = ('user__email', 'user__first_name', 'user__last_name', 'department_name')
    readonly_fields = ('created_at',)

    fieldsets = (
        ('Member', {
            'fields': ('user', 'graduation_year', 'department_name')
        }),
        ('Contribution', {
            'fields': ('offices_held', 'tasks_completed', 'proposals_passed')
        }),
        ('Message', {'fields': ('final_message',)}),
        ('Meta', {'fields': ('created_at',)}),
    )

    def user_col(self, obj):
        return obj.user.get_full_name() or obj.user.username
    user_col.short_description = 'Member'

    def offices_count(self, obj):
        return len(obj.offices_held or [])
    offices_count.short_description = 'Offices held'
    
    
from django.contrib import admin
from django.utils.html import format_html
from .models import SessionDocument


@admin.register(SessionDocument)
class SessionDocumentAdmin(admin.ModelAdmin):
    list_display = (
        'title', 'session_label', 'document_type',
        'is_published', 'is_featured', 'uploaded_at',
    )
    list_filter = (
        'document_type', 'session_year',
        'is_published', 'is_featured',
    )
    search_fields = ('title', 'description', 'session_label')
    list_editable = ('is_published', 'is_featured')
    raw_id_fields = ('uploaded_by',)
    readonly_fields = ('uploaded_at', 'updated_at')

    fieldsets = (
        ('Basic info', {
            'fields': ('title', 'description', 'document_type'),
        }),
        ('Session', {
            'fields': ('session_label', 'session_year'),
        }),
        ('Cover', {
            'fields': ('cover_image', 'external_cover_url'),
        }),
        ('Document file', {
            'fields': (
                'file', 'external_file_url', 'external_file_label',
            ),
            'description': (
                'Upload a file, or link to one (Google Drive, etc). '
                'External link takes priority if both are provided.'
            ),
        }),
        ('Visibility', {
            'fields': ('is_published', 'is_featured', 'order', 'uploaded_by'),
        }),
        ('Timestamps', {
            'fields': ('uploaded_at', 'updated_at'),
            'classes': ('collapse',),
        }),
    )    
    
    
    
    